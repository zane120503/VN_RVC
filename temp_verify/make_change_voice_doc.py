# -*- coding: utf-8 -*-
"""Sinh file Word HUONG_DAN_API_CHANGE_VOICE_AI.docx — API đổi giọng nam <-> nữ (/pitch_shift)."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = r"D:\Vietnamese-RVC\HUONG_DAN_API_CHANGE_VOICE_AI.docx"
BASE = "https://idolvoice.karaokeicool.vn"

doc = Document()

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
CODE_BG = "F2F2F2"
NOTE_BG = "FFF4E5"


def h1(text):
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        r.font.color.rgb = ACCENT
    return p


def h2(text):
    p = doc.add_heading(text, level=2)
    for r in p.runs:
        r.font.color.rgb = ACCENT
    return p


def para(text, bold=False, italic=False, keep_next=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if keep_next:
        p.paragraph_format.keep_with_next = True
    return p


def rich(*parts):
    """Đoạn văn ghép nhiều khúc; khúc bọc trong ** ** in đậm, `` in code."""
    p = doc.add_paragraph()
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(10)
            r.element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        else:
            p.add_run(part)
    return p


def bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    p.add_run(text)
    return p


def _shade(p, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    p._p.get_or_add_pPr().append(shd)


def code(text):
    """Khối code nền xám, Consolas; mỗi dòng 1 run + ngắt dòng mềm để giữ khối liền."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.keep_together = True   # khối code không bị cắt ngang qua 2 trang
    _shade(p, CODE_BG)
    lines = text.strip("\n").split("\n")
    for i, line in enumerate(lines):
        r = p.add_run(line)
        r.font.name = "Consolas"
        r.font.size = Pt(9)
        r.element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        if i < len(lines) - 1:
            r.add_break()
    return p


def note(text, label="Lưu ý: "):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.space_after = Pt(8)
    _shade(p, NOTE_BG)
    p.add_run(label).bold = True
    p.add_run(text)
    return p


def table(headers, rows, widths_cm=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        hdr[i].paragraphs[0].add_run(h).bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
            for pr in cells[i].paragraphs:
                for rr in pr.runs:
                    rr.font.size = Pt(10)
    for row in t.rows:   # 1 dòng bảng không bị cắt qua 2 trang
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)
    if widths_cm:
        for row in t.rows:
            for i, w in enumerate(widths_cm):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


# =====================================================================
# TIÊU ĐỀ
# =====================================================================
title = doc.add_heading("API ĐỔI GIỌNG NAM ↔ NỮ", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Change Voice — đổi giọng trên bản ghi âm karaoke của khách (Idol Voice)")
r.italic = True

table(["Thông tin", "Giá trị"], [
    ["Endpoint chính", "POST /pitch_shift"],
    ["Base URL", BASE],
    ["Xác thực", "Header X-API-Key: <API_KEY> (bắt buộc)"],
    ["Mô hình xử lý", "Bất đồng bộ: trả task_id ngay, xử lý theo hàng đợi 1 task/lúc"],
    ["Thời gian xử lý", "Khoảng 2–5 phút mỗi bài (không train model)"],
    ["Kết quả", "File MP3: giọng đã đổi + beat gốc"],
], widths_cm=[4, 12])

# =====================================================================
h1("1. Tổng quan")
para("API đổi giới tính giọng hát trên một bản ghi âm có sẵn của khách: giọng nam thành giọng nữ, "
     "hoặc giọng nữ thành giọng nam. Giao diện chỉ có 2 nút Nam / Nữ; mỗi nút gọi cùng một endpoint "
     "với giá trị semitones khác nhau.")
para("Quy trình xử lý bên trong một task:")
bullet("tách phần giọng hát ra khỏi nhạc nền (beat);", "1. Tách: ")
bullet("dịch cao độ riêng phần giọng đúng 1 quãng tám (±12 bán âm) và chỉnh formant (âm sắc khoang họng) "
       "để giọng nghe ra nam / nữ, không chỉ cao / thấp hơn;", "2. Đổi giọng: ")
bullet("tự cân âm lượng: giọng nhỏ hơn beat thì nâng lên bằng beat, giọng to hơn beat thì giữ nguyên;",
       "3. Cân âm lượng: ")
bullet("ghép giọng đã đổi với beat gốc thành file MP3.", "4. Ghép: ")
note("API không train model và không dùng AI học giọng, nên không cần dữ liệu giọng khách từ trước. "
     "Đổi lại, âm sắc vẫn là của người hát gốc được biến đổi, không phải một người khác giới thật hát.")

# =====================================================================
h1("2. Luồng tích hợp")
code("""
[Khách bấm nút Nam / Nữ trên UI]
        |
        v
POST /pitch_shift   (record_id, semitones)
        |
        v
<- Trả về ngay: { "status": "queued", "task_id": "..." }   <- LƯU task_id lại
        |
        |   (server xử lý 2–5 phút)
        v
Webhook POST về CRM: { "task_id", "status": "completed", "download_url" }
        |
        v
GET /download/{task_id}   -> file MP3 kết quả
""")
para("Không bắt buộc hỏi trạng thái liên tục: khi task xong, server tự gửi webhook về CRM (mục 5). "
     "Nếu cần hiển thị tiến độ thì gọi GET /status/{task_id} (mục 4).")

# =====================================================================
h1("3. POST /pitch_shift — Tạo task đổi giọng")
rich("Gửi dạng ", "`multipart/form-data`", ". Header bắt buộc: ", "`X-API-Key: <API_KEY>`", ".")

h2("3.1. Tham số dùng cho giao diện")
table(["Tham số", "Bắt buộc", "Kiểu", "Mô tả"], [
    ["record_id", "Có", "số nguyên", "ID bản ghi âm của khách trên hệ thống (id trong danh sách bản thu)"],
    ["semitones", "Có", "số nguyên", "12 = đổi ra giọng NỮ; -12 = đổi ra giọng NAM"],
], widths_cm=[3, 2, 2.5, 8.5])

h2("3.2. Ánh xạ 2 nút trên giao diện")
table(["Nút", "Gửi semitones", "Server tự áp dụng", "Dùng khi"], [
    ["Nữ", "12", "nâng +12 bán âm, formant ×1.18", "bản ghi âm giọng nam"],
    ["Nam", "-12", "hạ −12 bán âm, formant ×0.92", "bản ghi âm giọng nữ"],
], widths_cm=[2, 3, 6, 5])
note("Nên bấm nút ngược với giới tính người hát. Bản ghi giọng nam mà bấm Nam sẽ bị hạ thêm 1 quãng tám, "
     "nghe rất trầm và méo. API không tự nhận diện giới tính.")

h2("3.3. Tham số nâng cao (giao diện không cần gửi)")
table(["Tham số", "Mặc định", "Giới hạn", "Mô tả"], [
    ["formant_ratio", "0 (tự chọn)", "0.5 – 2.0",
     "0 = tự chọn theo hướng (×1.18 ra nữ, ×0.92 ra nam). 1 = chỉ đổi cao độ, không đổi formant. "
     "Số càng xa 1 thì càng nam/nữ hóa mạnh nhưng dễ rè hơn"],
    ["vocal_gain_db", "trống (tự cân)", "-12 – 24",
     "Bỏ trống = giọng nhỏ hơn beat thì nâng lên bằng beat, to hơn thì giữ nguyên. "
     "Truyền số = tự đặt mức tăng/giảm âm lượng giọng (dB)"],
    ["callback_url", "webhook CRM", "URL", "Ghi đè địa chỉ nhận webhook cho riêng task này"],
    ["semitones", "—", "-12 – 12, khác 0",
     "Ngoài ±12 có thể dùng mức khác (vd ±6). Lưu ý: không tròn ±12 thì giọng lệch tông với beat"],
], widths_cm=[3, 2.8, 2.4, 7.8])
rich("Để thử nghiệm, API còn nhận nguồn audio khác thay cho ", "`record_id`", ": upload file trực tiếp (",
     "`audio=@file.mp3`", ") hoặc bài trong danh mục (", "`target_song_id`",
     "). Giao diện khách chỉ dùng ", "`record_id`", ".")

h2("3.4. Ví dụ")
para("Bấm nút Nữ (bản ghi giọng nam):", bold=True, keep_next=True)
code(f"""
curl -X POST "{BASE}/pitch_shift" \\
  -H "X-API-Key: <API_KEY>" \\
  -F "record_id=319997" \\
  -F "semitones=12"
""")
para("Bấm nút Nam (bản ghi giọng nữ):", bold=True, keep_next=True)
code(f"""
curl -X POST "{BASE}/pitch_shift" \\
  -H "X-API-Key: <API_KEY>" \\
  -F "record_id=319997" \\
  -F "semitones=-12"
""")

h2("3.5. Response thành công (HTTP 200)")
code("""
{
  "status": "queued",
  "task_id": "74fb0984-bb29-496f-b973-1272e3083fa1",
  "semitones": 12,
  "song_name": "Karaoke Thiệp Hồng Sai Tên Tone Nam",
  "queue_size": 1,
  "eta_minutes": "1–3",
  "message": "Đang tách giọng và dịch +12 bán âm (thuần DSP, không train)."
}
""")
table(["Trường", "Ý nghĩa"], [
    ["task_id", "Mã task — BẮT BUỘC lưu lại để đối chiếu webhook, hỏi trạng thái, tải kết quả"],
    ["queue_size", "Số task đang xếp hàng (tính cả task này)"],
    ["eta_minutes", "Ước lượng thời gian khi hàng đợi trống; thực tế thường 2–5 phút với file dài"],
], widths_cm=[3.5, 12.5])

# =====================================================================
h1("4. GET /status/{task_id} — Xem trạng thái (tùy chọn)")
code(f"""
curl -H "X-API-Key: <API_KEY>" "{BASE}/status/74fb0984-bb29-496f-b973-1272e3083fa1"
""")
code("""
{
  "task_id": "74fb0984-bb29-496f-b973-1272e3083fa1",
  "status": "completed",
  "message": "Success",
  "result_path": "audios/record_319997_..._PITCH+12.mp3",
  "logs": "Đang tách giọng/beat...\\nTách xong. Đang dịch +12 bán âm + formant x1.18 (Praat)...\\nDịch xong. Đang ghép lại với beat gốc (giọng +0.0 dB)...\\n== HOÀN TẤT! ...",
  "kind": "pitch",
  "record_id": 319997
}
""")
table(["status", "Ý nghĩa"], [
    ["queued", "Đang chờ trong hàng đợi"],
    ["running", "Đang xử lý (xem logs để biết đang ở bước nào)"],
    ["completed", "Xong — tải kết quả qua /download/{task_id}"],
    ["failed", "Lỗi — lý do nằm trong message và logs"],
], widths_cm=[3.5, 12.5])

# =====================================================================
h1("5. Webhook báo kết quả về CRM")
para("Khi task kết thúc (thành công hoặc thất bại), server tự gửi:")
code("""
POST https://crm.icool.com.vn/api/ai-voice/webhook
Header: X-API-Key: <API_KEY>
Content-Type: application/json
""")
code(f"""
{{
  "event": "task_finished",
  "task_id": "74fb0984-bb29-496f-b973-1272e3083fa1",
  "kind": "pitch",
  "status": "completed",
  "message": "Success",
  "record_id": 319997,
  "customer_id": null,
  "song_id": null,
  "song_name": "Karaoke Thiệp Hồng Sai Tên Tone Nam",
  "download_url": "{BASE}/download/74fb0984-bb29-496f-b973-1272e3083fa1",
  "finished_at": "2026-09-30 03:15:42"
}}
""")
bullet("download_url = null khi status là failed.")
bullet("finished_at theo giờ UTC (CRM tự cộng 7 giờ khi hiển thị).")
bullet("Server thử gửi tối đa 3 lần, cách nhau 5 giây. Gửi không được thì chỉ ghi log; kết quả vẫn lấy được "
       "qua /status và /download.")
bullet("CRM đối chiếu theo task_id, nên phía gọi API phải lưu task_id nhận được ở bước 3. Task tạo bằng tay "
       "(không qua CRM) sẽ bị CRM trả 404 \"Không tìm thấy tác vụ\" — không ảnh hưởng kết quả.")

# =====================================================================
h1("6. GET /download/{task_id} — Tải file kết quả")
code(f"""
curl -H "X-API-Key: <API_KEY>" -o ket_qua.mp3 \\
  "{BASE}/download/74fb0984-bb29-496f-b973-1272e3083fa1"
""")
para("Trả file MP3 (audio/mpeg). Chỉ tải được khi status là completed; trước đó trả lỗi 400.")

# =====================================================================
h1("7. Mã lỗi")
table(["HTTP", "Khi nào", "Nội dung detail (ví dụ)"], [
    ["400", "semitones ngoài -12..12 hoặc bằng 0", "semitones phải trong [-12..12] và khác 0."],
    ["400", "Thiếu nguồn audio (không gửi record_id)", "Cần 1 nguồn audio: file audio, record_id hoặc target_song_id."],
    ["400", "formant_ratio / vocal_gain_db ngoài giới hạn", "formant_ratio trong [0.5..2.0] (hoặc 0 = tự chọn)."],
    ["400", "/download khi task chưa xong", "Task is not completed. Current status: running"],
    ["401", "Sai hoặc thiếu X-API-Key", "API key sai hoặc thiếu (header X-API-Key)."],
    ["404", "record_id không tồn tại", "Không tìm thấy bản ghi âm."],
    ["404", "task_id không tồn tại (/status, /download)", "Task ID not found"],
    ["502", "Không lấy được file ghi âm từ kho lưu trữ", "Không lấy được audio nguồn: ..."],
    ["503", "Server chưa kết nối được cơ sở dữ liệu", "DB chưa được cấu hình trên server."],
], widths_cm=[1.6, 6, 8.4])

# =====================================================================
h1("8. Ví dụ tích hợp đầy đủ (bash)")
code(f"""
KEY="<API_KEY>"
BASE="{BASE}"

# 1. Tạo task (nút Nữ)
TASK=$(curl -s -X POST "$BASE/pitch_shift" -H "X-API-Key: $KEY" \\
  -F "record_id=319997" -F "semitones=12" | jq -r .task_id)

# 2. (Tùy chọn) chờ xong — bình thường webhook sẽ báo, không cần vòng lặp này
until curl -s -H "X-API-Key: $KEY" "$BASE/status/$TASK" \\
  | jq -e '.status == "completed" or .status == "failed"' >/dev/null; do
  sleep 15
done

# 3. Tải kết quả
curl -s -H "X-API-Key: $KEY" -o ket_qua.mp3 "$BASE/download/$TASK"
""")

# =====================================================================
h1("9. Chất lượng và giới hạn")
bullet("Giọng đầu ra là người hát gốc được đổi cao độ và âm sắc, không phải một ca sĩ khác giới thật. "
       "Chuyển nữ ra nam thường nghe tự nhiên hơn nam ra nữ.")
bullet("Chất lượng phụ thuộc bản ghi âm. Bản thu mic karaoke có lẫn vang và beat thì phần tách giọng "
       "không sạch hoàn toàn; đổi formant mạnh dễ gây rè — đó là lý do hướng nam dùng ×0.92 thay vì mức mạnh hơn.")
bullet("Dùng ±12 (tròn 1 quãng tám) để giọng khớp tông với beat. Mức khác ±12 làm giọng lệch tông với nhạc.")
bullet("Bài nhiều người hát chung (song ca nam nữ, cả nhóm) sẽ đổi giọng tất cả cùng lúc.")

# =====================================================================
h1("10. Ghi chú vận hành")
bullet("Server xử lý tuần tự 1 task/lúc; nhiều request cùng lúc sẽ xếp hàng (xem queue_size).")
bullet("Danh sách task nằm trong bộ nhớ: khởi động lại server sẽ mất các task đang chờ hoặc đang chạy, "
       "và /status của chúng trả 404. Khi đó gọi lại POST /pitch_shift.")
bullet("File ghi âm gốc không bị thay đổi; mỗi lần đổi giọng tạo ra một file kết quả mới.")

doc.save(OUT)
print("Saved:", OUT)

# -*- coding: utf-8 -*-
"""Sinh file Word HUONG_DAN_API.docx từ nội dung tài liệu API."""
import docx
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ===== Style cơ bản =====
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

ACCENT = RGBColor(0x1F, 0x4E, 0x79)   # xanh đậm
CODE_BG = "F2F2F2"

def h1(text):
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        r.font.color.rgb = ACCENT
    return p

def h2(text):
    p = doc.add_heading(text, level=2)
    for r in p.runs:
        r.font.color.rgb = ACCENT
    return p

def h3(text):
    return doc.add_heading(text, level=3)

def para(text, bold=False, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    return p

def bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
    p.add_run(text)
    return p

def code(text):
    """Khối code: nền xám, chữ Consolas."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Cm(0.4)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), CODE_BG)
    p._p.get_or_add_pPr().append(shd)
    r = p.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(9)
    r.element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    return p

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        r = hdr[i].paragraphs[0].add_run(h)
        r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
            for pr in cells[i].paragraphs:
                for rr in pr.runs:
                    rr.font.size = Pt(10)
    doc.add_paragraph()
    return t

# =====================================================================
# TRANG BÌA / TIÊU ĐỀ
# =====================================================================
title = doc.add_heading("TÀI LIỆU API — IDOL VOICE", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Hệ thống đổi giọng ca sĩ bằng giọng khách hàng (Vietnamese-RVC Headless API)")
r.italic = True

para("")
table(["Thông tin", "Giá trị"], [
    ["Base URL", "https://idolvoice.karaokeicool.vn"],
    ["Xác thực", "Header X-API-Key: <key> — bắt buộc cho mọi endpoint trừ /health"],
    ["Mô hình xử lý", "Bất đồng bộ theo hàng đợi (trả task_id ngay, xử lý tuần tự 1 task/lúc)"],
    ["Swagger UI", "https://idolvoice.karaokeicool.vn/docs"],
])

# =====================================================================
h1("1. Luồng sử dụng chuẩn (theo khách hàng)")
code("""1. GET  /model/{customer_id}     -> khách đã có model chưa?
2. POST /train                   -> chưa có: train model từ file ghi âm (1 lần)
   GET  /status/{task_id}        -> poll đến khi completed
3. POST /convert                 -> đổi giọng bài hát bằng model của khách
   GET  /status/{task_id}        -> poll đến khi completed
4. GET  /download/{task_id}      -> tải file bài hát đã đổi giọng (.mp3)""")
bullet("Khách quay lại lần sau: bỏ qua bước 2, gọi thẳng /convert.")
bullet("Muốn train dữ liệu mới thay model cũ: /train với force_retrain=true.")

# =====================================================================
h1("2. Chi tiết các endpoint")

# ---- /health
h2("2.1. GET /health — Kiểm tra server")
para("Không cần API key.")
code("curl https://idolvoice.karaokeicool.vn/health")
para("Response 200:")
code('{"status": "ok", "mode": "headless-async", "active_tasks": 3}')

# ---- /train
h2("2.2. POST /train — Train model giọng theo khách hàng")
para("Upload file ghi âm của khách → hệ thống tách giọng, huấn luyện model RVC và lưu theo customer_id (đăng ký vào DB bảng rvc_customer_models).")
para("Content-Type: multipart/form-data", italic=True)
table(["Tham số", "Kiểu", "Bắt buộc", "Mặc định", "Mô tả"], [
    ["customer_id", "text", "✓", "—", "Mã khách hàng (chữ/số/gạch)"],
    ["training_files", "file (nhiều)", "✓", "—", "File ghi âm giọng khách (wav/mp3...). Lặp lại field để gửi nhiều file"],
    ["epochs", "int", "", "150", "Số vòng huấn luyện (100–300 khuyến nghị)"],
    ["force_retrain", "bool", "", "false", "true = xóa model cũ, train dữ liệu mới thay thế"],
])
code("""curl -X POST https://idolvoice.karaokeicool.vn/train \\
  -H "X-API-Key: YOUR_API_KEY" \\
  -F "customer_id=KH001" \\
  -F "epochs=150" \\
  -F "force_retrain=false" \\
  -F "training_files=@ghi_am_1.wav" \\
  -F "training_files=@ghi_am_2.wav\"""")
para("Response — khách mới (bắt đầu train):")
code("""{
  "status": "queued",
  "task_id": "55d9ca4c-d023-4abe-a680-c6d85cd9f379",
  "customer_id": "KH001",
  "model_name": "cus_KH001",
  "message": "Đã nhận 2 file ghi âm. Bắt đầu train model (vị trí hàng đợi: 1).",
  "queue_size": 1
}""")
para("Response — khách đã có model (không train lại):")
code("""{
  "status": "exists",
  "customer_id": "KH001",
  "model_name": "cus_KH001",
  "model_file": "cus_KH001_150e_....pth",
  "message": "Khách hàng đã có model. Gửi force_retrain=true nếu muốn train dữ liệu mới thay thế."
}""")
para("Yêu cầu dữ liệu train:", bold=True)
bullet("Tổng thời lượng giọng nói thực tế (đã trừ khoảng lặng) tối thiểu 60 giây — ít hơn task sẽ failed.")
bullet("Nên gửi 10–30 phút ghi âm, rõ, ít tạp âm để chất lượng tốt.")
bullet("Thời gian train: ~20–60 phút tùy dữ liệu và epochs.")

# ---- /model
h2("2.3. GET /model/{customer_id} — Kiểm tra model của khách")
code('curl -H "X-API-Key: YOUR_API_KEY" https://idolvoice.karaokeicool.vn/model/KH001')
para("Response 200:")
code("""{
  "customer_id": "KH001",
  "trained": true,
  "model_name": "cus_KH001",
  "model_file": "cus_KH001_150e_....pth",
  "index_file": "added_IVF..._cus_KH001_v2.index",
  "db_record": { "epochs": 150, "trained_at": "2026-07-09 09:15:00", ... }
}""")
para("trained: false → cần gọi /train trước khi /convert.")

# ---- /convert
h2("2.4. POST /convert — Đổi giọng bài hát bằng model của khách")
para("Dùng model đã train của customer_id để thay giọng ca sĩ trong bài hát, tự ghép lại beat.")
para("Content-Type: multipart/form-data", italic=True)
table(["Tham số", "Kiểu", "Bắt buộc", "Mặc định", "Mô tả"], [
    ["customer_id", "text", "✓", "—", "Mã khách (phải có model — nếu chưa: lỗi 404)"],
    ["target_song_id", "text", "chọn 1", "—", "ID bài hát trong hệ thống — server tự lấy file audio"],
    ["target_song", "file", "chọn 1", "—", "HOẶC upload file bài hát trực tiếp (mp3/wav...)"],
    ["pitch_shift", "int", "", "0", "Dịch cao độ nửa cung: cùng giới 0; nam→nữ +12; nữ→nam -12"],
])
para("Cách 1 — theo id bài hát (khuyên dùng):")
code("""curl -X POST https://idolvoice.karaokeicool.vn/convert \\
  -H "X-API-Key: YOUR_API_KEY" \\
  -F "customer_id=KH001" \\
  -F "target_song_id=103691" \\
  -F "pitch_shift=0\"""")
para("Cách 2 — upload file bài hát:")
code("""curl -X POST https://idolvoice.karaokeicool.vn/convert \\
  -H "X-API-Key: YOUR_API_KEY" \\
  -F "customer_id=KH001" \\
  -F "target_song=@bai_hat.mp3" \\
  -F "pitch_shift=0\"""")
para("Response 200:")
code("""{
  "status": "queued",
  "task_id": "a1b2c3d4-...",
  "customer_id": "KH001",
  "model_name": "cus_KH001",
  "message": "Bắt đầu đổi giọng bằng model của khách (vị trí hàng đợi: 1).",
  "queue_size": 1
}""")
para("Lưu ý target_song_id:", bold=True)
bullet("Chỉ dùng được bài đã xuất bản/đồng bộ lên media server. Bài chưa sync → lỗi 502 (\"bài chưa xuất bản (version v/0)\").")
bullet("Thời gian convert: ~2–5 phút/bài.")

# ---- /status
h2("2.5. GET /status/{task_id} — Trạng thái task")
code('curl -H "X-API-Key: YOUR_API_KEY" https://idolvoice.karaokeicool.vn/status/55d9ca4c-...')
para("Response 200:")
code("""{
  "task_id": "55d9ca4c-...",
  "status": "running",
  "message": "Processing...",
  "result_path": null,
  "logs": "== BẮT ĐẦU HUẤN LUYỆN (150 epochs) ==\\nĐang tiền xử lý...",
  "kind": "train",
  "customer_id": "KH001"
}""")
para("Vòng đời status:  queued → running → completed | failed", bold=True)
bullet("logs chứa 3000 ký tự log gần nhất — dùng hiển thị tiến độ hoặc debug khi failed.")
bullet("Khuyến nghị poll mỗi 10–15 giây.")
bullet("Task lưu trong RAM — restart server sẽ mất task_id cũ (model đã train KHÔNG mất).")

# ---- /download
h2("2.6. GET /download/{task_id} — Tải kết quả")
para("Chỉ dùng được khi task completed (chưa xong → 400).")
code('curl -H "X-API-Key: YOUR_API_KEY" https://idolvoice.karaokeicool.vn/download/a1b2c3d4-... -o ket_qua.mp3')
bullet("Task convert: trả file bài hát đã đổi giọng (.mp3).")
bullet("Task train: trả file model .pth (thường không cần tải — model đã lưu trên server).")
bullet("File kết quả giữ trên server 10 ngày rồi tự xóa.")

# ---- /songs
h2("2.7. GET /songs — Danh sách / tìm kiếm bài hát")
para("Trả về danh sách bài hát trong hệ thống (chỉ các bài có media), có phân trang.")
table(["Tham số (query)", "Kiểu", "Mặc định", "Mô tả"], [
    ["q", "text", "—", "Từ khóa tìm theo tên bài (có dấu hoặc không dấu đều được)"],
    ["limit", "int", "50", "Số bài mỗi trang (tối đa 200)"],
    ["offset", "int", "0", "Vị trí bắt đầu (phân trang)"],
])
code("""# 50 bài mới nhất
curl -H "X-API-Key: YOUR_API_KEY" "https://idolvoice.karaokeicool.vn/songs"

# Tìm bài theo tên (không dấu cũng được)
curl -H "X-API-Key: YOUR_API_KEY" "https://idolvoice.karaokeicool.vn/songs?q=hen%20yeu&limit=20"

# Trang tiếp theo
curl -H "X-API-Key: YOUR_API_KEY" "https://idolvoice.karaokeicool.vn/songs?limit=50&offset=50\"""")
para("Response 200:")
code("""{
  "total": 2,
  "limit": 20,
  "offset": 0,
  "count": 2,
  "songs": [
    {"id": 107929, "name": "HẸN YÊU", "duration": 301, "version": 26052200},
    {"id": 104717, "name": "HẸN YÊU", "duration": 328, "version": 25102800}
  ]
}""")
para("⚠️ Bài thuộc build mới nhất có thể chưa đồng bộ lên media server. Trước khi /convert, nên xác nhận bằng /check_song/{id}.", bold=True)

# ---- /check_song
h2("2.8. GET /check_song/{song_id} — Kiểm tra bài có sẵn để convert")
code('curl -H "X-API-Key: YOUR_API_KEY" https://idolvoice.karaokeicool.vn/check_song/107929')
para("Response — bài dùng được:")
code('{"song_id": "107929", "available": true, "size_mb": 9.2, "reason": null}')
para("Response — bài chưa đồng bộ:")
code("""{"song_id": "108578", "available": false, "size_mb": null,
 "reason": "Bài chưa xuất bản/đồng bộ lên media server (version v/0)"}""")
para("Luồng khuyến nghị cho app: /songs?q=... cho khách chọn bài → /check_song/{id} xác nhận → /convert.", italic=True)

# ---- legacy
h2("2.9. API cũ (giữ để tương thích)")
h3("POST /run_upload — Train + convert trong 1 lần gọi")
code("""curl -X POST https://idolvoice.karaokeicool.vn/run_upload \\
  -H "X-API-Key: YOUR_API_KEY" \\
  -F "model_name=TenModel" \\
  -F "epochs=150" -F "pitch_shift=0" -F "force_retrain=false" \\
  -F "target_song_id=103666" \\
  -F "training_files=@giong.wav\"""")
h3("POST /run — Như trên nhưng dùng đường dẫn file có sẵn trên server (JSON)")
code("""curl -X POST https://idolvoice.karaokeicool.vn/run \\
  -H "X-API-Key: YOUR_API_KEY" \\
  -H "Content-Type: application/json" \\
  -d '{"training_files":["/app/dataset/giong.wav"],"target_song_path":"/app/audios/bai.mp3",
       "model_name":"TenModel","epochs":150,"pitch_shift":0,"force_retrain":false}'""")
para("Khuyến nghị dùng bộ API mới /train + /convert — tách bạch, tái sử dụng model theo khách hàng.", italic=True)

# =====================================================================
h1("3. Bảng tổng hợp endpoint")
table(["#", "Method", "Endpoint", "Chức năng", "API Key"], [
    ["1", "GET", "/health", "Kiểm tra server sống", "Không"],
    ["2", "POST", "/train", "Train model giọng theo khách hàng", "Có"],
    ["3", "GET", "/model/{customer_id}", "Kiểm tra model của khách", "Có"],
    ["4", "POST", "/convert", "Đổi giọng bài hát bằng model của khách", "Có"],
    ["5", "GET", "/status/{task_id}", "Trạng thái + log task", "Có"],
    ["6", "GET", "/download/{task_id}", "Tải file kết quả", "Có"],
    ["7", "GET", "/songs", "Danh sách / tìm kiếm bài hát", "Có"],
    ["8", "GET", "/check_song/{song_id}", "Kiểm tra bài có sẵn để convert", "Có"],
    ["9", "POST", "/run_upload", "(Cũ) Train + convert 1 lần", "Có"],
    ["10", "POST", "/run", "(Cũ) Như trên, input là path trên server", "Có"],
])

# =====================================================================
h1("4. Mã lỗi thường gặp")
table(["HTTP", "Ý nghĩa", "Cách xử lý"], [
    ["400", "Thiếu tham số / task chưa completed khi download", "Kiểm tra request"],
    ["401", "Sai hoặc thiếu X-API-Key", "Gửi đúng header"],
    ["404", "Khách chưa có model / task_id không tồn tại / bài hát không có media", "Train trước; kiểm tra id"],
    ["502", "Không lấy được audio theo target_song_id (bài chưa xuất bản v/0, media server lỗi)", "Dùng id bài đã xuất bản hoặc upload file"],
    ["500", "Lỗi xử lý nội bộ", "Xem logs trong /status, báo quản trị"],
])
para("Task failed (qua /status): nguyên nhân phổ biến — dữ liệu train < 60 giây giọng thực tế, file âm thanh hỏng, hết VRAM. Đọc trường logs để biết chi tiết.")

# =====================================================================
h1("5. Ví dụ tích hợp — luồng đầy đủ (bash)")
code("""BASE="https://idolvoice.karaokeicool.vn"
KEY="YOUR_API_KEY"
CUS="KH001"

# 1) Khách đã có model chưa?
TRAINED=$(curl -s -H "X-API-Key: $KEY" "$BASE/model/$CUS" | grep -o '"trained":[a-z]*' | cut -d: -f2)

# 2) Chưa có -> train
if [ "$TRAINED" != "true" ]; then
  TASK=$(curl -s -X POST "$BASE/train" -H "X-API-Key: $KEY" \\
    -F "customer_id=$CUS" -F "epochs=150" \\
    -F "training_files=@ghi_am.wav" | grep -o '"task_id":"[^"]*"' | cut -d'"' -f4)
  while :; do
    S=$(curl -s -H "X-API-Key: $KEY" "$BASE/status/$TASK" | grep -o '"status":"[^"]*"' | head -1 | cut -d'"' -f4)
    echo "train: $S"; [ "$S" = "completed" ] && break
    [ "$S" = "failed" ] && exit 1
    sleep 15
  done
fi

# 3) Convert bài hát
TASK=$(curl -s -X POST "$BASE/convert" -H "X-API-Key: $KEY" \\
  -F "customer_id=$CUS" -F "target_song_id=103691" -F "pitch_shift=0" \\
  | grep -o '"task_id":"[^"]*"' | cut -d'"' -f4)
while :; do
  S=$(curl -s -H "X-API-Key: $KEY" "$BASE/status/$TASK" | grep -o '"status":"[^"]*"' | head -1 | cut -d'"' -f4)
  echo "convert: $S"; [ "$S" = "completed" ] && break
  [ "$S" = "failed" ] && exit 1
  sleep 15
done

# 4) Tải kết quả
curl -H "X-API-Key: $KEY" "$BASE/download/$TASK" -o ket_qua.mp3""")

# =====================================================================
h1("6. Ghi chú vận hành")
bullet("pitch_shift (nửa cung): cùng giới tính 0; model nam hát bài ca sĩ nữ -12; model nữ hát bài ca sĩ nam +12; lệch nhẹ thử ±3..6.")
bullet("Model theo khách được giữ vĩnh viễn (assets/weights + DB) — convert các lần sau không cần train lại.")
bullet("File upload input tự xóa sau khi task xong; file kết quả trong audios/ tự xóa sau 10 ngày.")
bullet("Server xử lý tuần tự — nhiều request cùng lúc sẽ xếp hàng (xem queue_size trong response).")
bullet("Task registry nằm trong RAM: restart server làm mất task_id đang theo dõi (model/kết quả trên đĩa không mất).")

doc.save(r"d:\\Vietnamese-RVC\\HUONG_DAN_API.docx")
print("Saved: d:\\Vietnamese-RVC\\HUONG_DAN_API.docx")
